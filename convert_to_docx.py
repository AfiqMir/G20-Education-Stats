import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "D3D3D3")}"/>'
            tcBorders.append(parse_xml(tag))
        else:
            tag = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_formatted_runs(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            link_text = part[1:part.index('](')]
            link_url = part[part.index('](')+2:-1]
            run = paragraph.add_run(link_text)
            run.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
            run.underline = True
        else:
            paragraph.add_run(part)

def create_full_report_docx(md_filepath='laporan.md', output_docx='Laporan_Prapemrosesan_EdStats.docx'):
    doc = Document()
    
    # Page setup: Standard A4, 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
    # Default Normal Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(4)

    with open(md_filepath, 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.split('\n')
    i = 0

    def add_styled_heading(text, level):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        if level == 1:
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Dark Navy
        elif level == 2:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        elif level == 3:
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(2)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        return p

    def add_image_if_exists(img_filename, caption_text, width=Inches(6.2)):
        if os.path.exists(img_filename):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(img_filename, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            run_cap = p_cap.add_run(f'Gambar: {caption_text}')
            run_cap.italic = True
            run_cap.font.size = Pt(9)
            run_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    inserted_images = set()

    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.strip()
        
        if not line:
            i += 1
            continue
            
        # Title (Main Document Title)
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(17)
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            i += 1
            continue
            
        # Headings
        if line.startswith('### '):
            add_styled_heading(line[4:], 3)
            i += 1
            continue
        elif line.startswith('## '):
            add_styled_heading(line[3:], 2)
            i += 1
            continue
            
        # Horizontal Rule (Divider)
        if line.startswith('---'):
            i += 1
            continue
            
        # Blockquotes / Callout notes
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run_tag = p.add_run("💡 ")
            add_formatted_runs(p, line[2:].replace('> ', ''))
            i += 1
            continue

        # Markdown Tables
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            header_row = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            num_cols = len(header_row)
            
            data_rows = []
            for t_line in table_lines[2:]: # skip separator line
                cols = [c.strip() for c in t_line.split('|')[1:-1]]
                if len(cols) == num_cols:
                    data_rows.append(cols)
                    
            table = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            
            # Style header row
            for col_idx, col_name in enumerate(header_row):
                cell = table.cell(0, col_idx)
                cell.text = ""
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                add_formatted_runs(p, col_name)
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell, 80, 80, 100, 100)
                set_cell_border(cell, top={"sz": 4, "color": "1B365D"}, bottom={"sz": 6, "color": "1B365D"})
                
            # Style data rows
            for row_idx, r_data in enumerate(data_rows):
                fill = "F7F9FC" if row_idx % 2 == 1 else "FFFFFF"
                for col_idx, val in enumerate(r_data):
                    cell = table.cell(row_idx + 1, col_idx)
                    cell.text = ""
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    add_formatted_runs(p, val)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                    set_cell_background(cell, fill)
                    set_cell_margins(cell, 60, 60, 80, 80)
                    set_cell_border(cell, bottom={"sz": 4, "color": "E2E8F0"})
                    
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(4)
            continue
            
        # Bullet list items with indentation levels
        leading_spaces = len(raw_line) - len(raw_line.lstrip())
        
        if line.startswith('- ') or line.startswith('* '):
            content = line[2:]
            p = doc.add_paragraph()
            # Determine bullet level based on leading spaces
            if leading_spaces >= 6:
                p.paragraph_format.left_indent = Inches(0.65)
                run_b = p.add_run("▪  ")
            elif leading_spaces >= 3:
                p.paragraph_format.left_indent = Inches(0.45)
                run_b = p.add_run("◦  ")
            else:
                p.paragraph_format.left_indent = Inches(0.25)
                run_b = p.add_run("•  ")
            
            run_b.bold = True
            run_b.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, content)
            
            # Check context triggers for images
            if "Tingkat Putus Sekolah" in content and 'boxplot' not in inserted_images and "Langkah 4" in md_text:
                add_image_if_exists('boxplot_outlier_detection.png', 'Distribusi Boxplot dan Deteksi Outlier 6 Indikator Pendidikan Dasar G20', Inches(6.0))
                inserted_images.add('boxplot')

            i += 1
            continue
            
        # Numbered list items
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line)
        if num_match:
            num_str = num_match.group(1)
            content = num_match.group(2)
            p = doc.add_paragraph()
            if leading_spaces >= 3:
                p.paragraph_format.left_indent = Inches(0.45)
            else:
                p.paragraph_format.left_indent = Inches(0.25)
            run_num = p.add_run(f"{num_str}.  ")
            run_num.bold = True
            run_num.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, content)
            
            # Check list item triggers for figures
            if "Posisi Benchmarking Indonesia" in content and 'tren' not in inserted_images:
                add_image_if_exists('visualisasi_tren_pendidikan_g20.png', 'Tren Angka Partisipasi Murni (NER) 19 Negara Anggota G20 (2000–2017)', Inches(6.2))
                inserted_images.add('tren')
            elif "Paritas Gender di G20" in content and 'parity' not in inserted_images:
                add_image_if_exists('gender_parity_analysis.png', 'Analisis Paritas Gender Partisipasi Siswa Laki-laki vs. Perempuan terhadap Garis 1:1', Inches(5.2))
                inserted_images.add('parity')
            elif "Peta Panas Evolusi Pendidikan" in content and 'heatmap' not in inserted_images:
                add_image_if_exists('temporal_heatmap_g20.png', 'Peta Panas (Heatmap) Evolusi Partisipasi Murni 19 Negara G20 (2000–2017)', Inches(6.2))
                inserted_images.add('heatmap')
            elif "Studi Kasus Efektivitas Imputasi" in content and 'case_study' not in inserted_images:
                add_image_if_exists('before_after_imputation_case_study.png', 'Studi Kasus Before vs. After Imputasi Deret Waktu (Indonesia & India)', Inches(6.2))
                inserted_images.add('case_study')

            i += 1
            continue
            
        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_runs(p, line)
        
        # Check context triggers in paragraphs
        if "Total Informasi Tersimpan" in line and 'pca' not in inserted_images:
            add_image_if_exists('pca_scree_plot.png', 'Scree Plot dan Retensi Kumulatif Variansi PCA (84.47%)', Inches(5.5))
            add_image_if_exists('pca_2d_projection.png', 'Proyeksi Sebaran 2D PCA Space Negara Anggota G20', Inches(5.8))
            inserted_images.add('pca')
        elif "Ringkasan Evaluasi Kuantitatif" in line and 'before_after' not in inserted_images:
            add_image_if_exists('before_after_comparison.png', 'Evaluasi Sebelum vs. Sesudah Preprocessing: Missing Rate 3 Tahap & Efisiensi Penyimpanan', Inches(6.0))
            inserted_images.add('before_after')

        i += 1

    # End check
    if 'before_after' not in inserted_images:
        add_image_if_exists('before_after_comparison.png', 'Evaluasi Sebelum vs. Sesudah Preprocessing: Missing Rate 3 Tahap & Efisiensi Penyimpanan', Inches(6.0))

    doc.save(output_docx)
    print(f'Dokumen Word berhasil dibuat: {output_docx} (Ukuran: {os.path.getsize(output_docx)/1024:.1f} KB)')

if __name__ == '__main__':
    create_full_report_docx()

