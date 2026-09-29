import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_pptx='Presentasi_Prapemrosesan_EdStats.pptx'):
    prs = Presentation()
    # 16:9 Widescreen standard dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # blank layout

    # Professional Color Palette
    COLOR_PRIMARY = RGBColor(0x1B, 0x36, 0x5D)    # Dark Navy
    COLOR_SECONDARY = RGBColor(0x02, 0x84, 0xC7)  # Vibrant Blue
    COLOR_ACCENT = RGBColor(0x0D, 0x94, 0x88)     # Teal / Green accent
    COLOR_DARK_TEXT = RGBColor(0x1E, 0x29, 0x3B)  # Slate 800
    COLOR_MUTED_TEXT = RGBColor(0x64, 0x74, 0x8B) # Slate 500
    COLOR_BG_CARD = RGBColor(0xF8, 0xFA, 0xFC)    # Light Slate Card
    COLOR_CARD_BORDER = RGBColor(0xE2, 0xE8, 0xF0)# Border Slate
    COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)      # Pure White

    def add_header(slide, title_text, category_text="REKAYASA DATA (DATA ENGINEERING) - UGM"):
        # Category / Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_SECONDARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

    def add_card(slide, left, top, width, height, bg_color=COLOR_BG_CARD, border_color=COLOR_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    
    # Background accent card
    add_card(slide1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=COLOR_PRIMARY, border_color=None)
    
    # Inner light card
    add_card(slide1, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER)

    t_box = slide1.shapes.add_textbox(Inches(1.6), Inches(1.6), Inches(10.1), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "PROYEK REKAYASA DATA"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY
    p.space_after = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Alur Kerja Data Preprocessing pada World Bank EdStats"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Integrasi 5 Tabel Lengkap & Benchmarking Pendidikan Dasar 19 Negara G20 (2000–2017)"
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_MUTED_TEXT

    # Author Box
    auth_card = add_card(slide1, Inches(1.6), Inches(4.2), Inches(10.1), Inches(1.6), bg_color=COLOR_BG_CARD, border_color=COLOR_CARD_BORDER)
    auth_box = slide1.shapes.add_textbox(Inches(1.8), Inches(4.3), Inches(9.7), Inches(1.4))
    tf_auth = auth_box.text_frame
    tf_auth.word_wrap = True
    
    p = tf_auth.paragraphs[0]
    p.text = "MATA KULIAH:"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY
    
    p = tf_auth.add_paragraph()
    p.text = "Rekayasa Data (Data Engineering)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    p = tf_auth.add_paragraph()
    p.text = "Departemen Teknik Elektro dan Teknologi Informasi, Fakultas Teknik UGM"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK_TEXT

    # =========================================================================
    # SLIDE 2: LATAR BELAKANG & TANTANGAN DATA MENTAH
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Latar Belakang & Karakteristik Data Mentah")

    # Card 1: Profil Data
    add_card(slide2, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    b1 = slide2.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(4.9))
    tf = b1.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Dataset World Bank EdStats"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    points = [
        ("Karakteristik Skala Besar:", "886.930 baris x 70 kolom (~326,4 MB) dengan alokasi memori ~740 MB."),
        ("Struktur Wide Format:", "Setiap tahun dari 1970 hingga 2100 direpresentasikan dalam kolom terpisah."),
        ("Ekosistem 5 Berkas Terpisah:", "Terdiri dari EdStatsData (fakta), EdStatsCountry, EdStatsSeries, EdStatsFootNote, dan EdStatsCountry-Series."),
        ("Fokus Jenjang:", "Pendidikan Dasar (Primary Education / ISCED Level 1) untuk kohort usia 7–12 tahun.")
    ]
    for title, desc in points:
        p = tf.add_paragraph()
        p.text = f"• {title} "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        p.space_after = Pt(8)

    # Card 2: 3 Masalah Kualitas
    add_card(slide2, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    b2 = slide2.shapes.add_textbox(Inches(7.2), Inches(1.7), Inches(5.0), Inches(4.9))
    tf2 = b2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "Tantangan Utama Kualitas Data"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    issues = [
        ("1. Ekstremitas Nilai Kosong (86,10% Missing):", "Sebagian besar sel kosong karena proyeksi jangka panjang tidak terisi."),
        ("2. Noise & Kolom Anomali:", "Terdapat kolom kosong tanpa nama (Unnamed: 69) dan rentang proyeksi teoritis masa depan (>2017)."),
        ("3. Konflik Format & Ketiadaan Relasi:", "Format tahun string (YR2000) dan belum terintegrasinya catatan metodologi survei/sensus ke tabel fakta.")
    ]
    for title, desc in issues:
        p = tf2.add_paragraph()
        p.text = f"• {title}\n  "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 3: ALUR DATA PREPROCESSING (7 LANGKAH)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Alur Kerja End-to-End Data Preprocessing (7 Langkah)")

    steps = [
        ("Langkah 1", "Pemuatan 5 Berkas Data", "Pemeriksaan dimensi mentah & evaluasi missing rate"),
        ("Langkah 2", "Strategi Reduksi Data", "Filter 19 Negara G20, kurun 2000–2017, & 6 indikator inti"),
        ("Langkah 3", "Integrasi 5 Tabel Lengkap", "Resolusi format YR2000, Star Schema, & Uji Chi-Square"),
        ("Langkah 4", "Pembersihan & Imputasi", "Pivoting Tidy Panel, Imputasi 3-Tier, & Deteksi Outlier"),
        ("Langkah 5", "Reduksi Dimensi (PCA)", "Standarisasi Z-Score & ekstraksi 2 Komponen Utama (84,5%)"),
        ("Langkah 6", "Diskretisasi Data", "Binning interval tingkat kelulusan & rasio beban guru"),
        ("Langkah 7", "Evaluasi & Visualisasi", "Validasi Before vs After & analisis tren 19 negara G20")
    ]

    for idx, (s_num, s_title, s_desc) in enumerate(steps):
        row = idx // 4
        col = idx % 4
        x = Inches(0.8 + col * 2.95)
        y = Inches(1.5 + row * 2.6)
        w = Inches(2.8)
        h = Inches(2.3) if row == 0 else Inches(2.3)
        
        # Highlight last step or normal card
        bg = RGBColor(0xED, 0xF8, 0xFD) if idx == 6 else COLOR_BG_CARD
        bdr = COLOR_SECONDARY if idx == 6 else COLOR_CARD_BORDER
        add_card(slide3, x, y, w, h, bg_color=bg, border_color=bdr)
        
        tb = slide3.shapes.add_textbox(x + Inches(0.15), y + Inches(0.15), w - Inches(0.3), h - Inches(0.3))
        tf_s = tb.text_frame
        tf_s.word_wrap = True
        
        p = tf_s.paragraphs[0]
        p.text = s_num
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_SECONDARY
        
        p = tf_s.add_paragraph()
        p.text = s_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.space_after = Pt(4)

        p = tf_s.add_paragraph()
        p.text = s_desc
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_DARK_TEXT

    # =========================================================================
    # SLIDE 4: LANGKAH 2 & 3: REDUKSI DATA & INTEGRASI 5 TABEL
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "Langkah 2 & 3: Reduksi Data dan Integrasi 5 Berkas")

    add_card(slide4, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    tb_l2 = slide4.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(4.9))
    tf = tb_l2.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Langkah 2: Strategi Reduksi 3 Dimensi"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)

    p_list = [
        ("1. Dimensi Geografis:", "19 Negara Berdaulat G20 (representasi kekuatan ekonomi dunia & negara berkembang)."),
        ("2. Dimensi Temporal:", "Rentang 2000–2017 kontinu (mengeliminasi kolom proyeksi masa depan)."),
        ("3. Fitur Indikator (Net Enrolment Focus):", "6 Indikator Pendidikan Dasar: Partisipasi Murni (Total, Perempuan, Laki-laki), Kelulusan Dasar, Rasio Guru, dan Putus Sekolah."),
        ("💡 Keuntungan Domain:", "Seluruh indikator partisipasi murni menggunakan basis kohort usia yang sama persis (7–12 tahun), berada pada skala logis 0%–100%.")
    ]
    for t, d in p_list:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        run = p.add_run()
        run.text = d
        run.font.bold = False
        p.space_after = Pt(6)

    add_card(slide4, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tb_l3 = slide4.shapes.add_textbox(Inches(7.2), Inches(1.7), Inches(5.0), Inches(4.9))
    tf = tb_l3.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Langkah 3: Integrasi 5 Tabel & Resolusi Konflik"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)

    p_list2 = [
        ("Penyatuan Skema Bintang:", "Left join bertingkat antara tabel fakta dengan Country, Series, FootNote, dan Country-Series."),
        ("Resolusi Konflik Format:", "Pembersihan regex format string 'YR2000' menjadi integer 2000."),
        ("Ekstraksi Atribut Silsilah (Provenance):", "Membentuk fitur biner 'Is_Estimated' (membedakan data sensus lapangan langsung vs estimasi model)."),
        ("Uji Chi-Square Dependensi Redundansi:", "Uji antara Income Group dan Region menghasilkan Chi-Square = 557,33 (p-value = 8,80e-107 < 0,05), membuktikan dependensi regional sangat kuat.")
    ]
    for t, d in p_list2:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        run = p.add_run()
        run.text = d
        run.font.bold = False
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 5: LANGKAH 4: PEMBERSIHAN DATA & IMPUTASI 3-TIER
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Langkah 4: Pembersihan Data, Format Panel, & Imputasi 3-Tier")

    # Left: Text
    add_card(slide5, Inches(0.8), Inches(1.5), Inches(4.8), Inches(5.3))
    tb = slide5.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.4), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Hierarchical 3-Tier Imputation"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(6)

    imp_points = [
        ("Format Tidy Panel:", "Data di-unpivot dan di-pivot menjadi 342 baris observasi (19 Negara x 18 Tahun)."),
        ("Tingkat 1 (Interpolasi Linier Temporal):", "Mengisi titik waktu kosong di antara tahun yang ada untuk menjaga kontinuitas tren deret waktu."),
        ("Tingkat 2 (Median Kelompok Pendapatan):", "Mengisi gap batas awal/akhir berdasarkan median kelompok pendapatan negara sebaya (Peer Income Group)."),
        ("Tingkat 3 (Median Global G20):", "Fallback nilai median untuk menjamin 0% missing value."),
        ("Hasil Evaluasi:", "Missing values tuntas dari 46,54% menjadi 0,00% tanpa distorsi distribusi.")
    ]
    for t, d in imp_points:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK_TEXT
        run = p.add_run()
        run.text = d
        run.font.bold = False
        p.space_after = Pt(5)

    # Right: Boxplot Outlier Image
    add_card(slide5, Inches(5.8), Inches(1.5), Inches(6.7), Inches(5.3), bg_color=COLOR_WHITE)
    if os.path.exists('boxplot_outlier_detection.png'):
        slide5.shapes.add_picture('boxplot_outlier_detection.png', Inches(6.0), Inches(1.8), width=Inches(6.3))

    # =========================================================================
    # SLIDE 6: LANGKAH 5: REDUKSI DIMENSI DENGAN PCA
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "Langkah 5: Reduksi Dimensi dengan Principal Component Analysis (PCA)")

    # Left: Scree Plot Image
    add_card(slide6, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.2), bg_color=COLOR_WHITE)
    if os.path.exists('pca_scree_plot.png'):
        slide6.shapes.add_picture('pca_scree_plot.png', Inches(0.95), Inches(1.65), width=Inches(5.3))

    # Right: 2D Scatter Image
    add_card(slide6, Inches(6.9), Inches(1.5), Inches(5.6), Inches(4.2), bg_color=COLOR_WHITE)
    if os.path.exists('pca_2d_projection.png'):
        slide6.shapes.add_picture('pca_2d_projection.png', Inches(7.05), Inches(1.65), width=Inches(5.3))

    # Bottom: Summary Cards
    add_card(slide6, Inches(0.8), Inches(5.85), Inches(11.7), Inches(1.2), bg_color=COLOR_PRIMARY, border_color=None)
    tb_pca = slide6.shapes.add_textbox(Inches(1.0), Inches(5.9), Inches(11.3), Inches(1.0))
    tf = tb_pca.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Hasil Kunci Reduksi Dimensi PCA (Standarisasi Z-Score):"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY

    p = tf.add_paragraph()
    p.text = "• PC1 (70,71% Variansi): Sumbu Akses Partisipasi & Kelulusan   |   • PC2 (13,76% Variansi): Sumbu Rasio Beban Guru\n• Total Retensi Informasi (2 PC): 84,47% (Melebihi threshold standar 80% dan memisahkan klaster negara maju vs berkembang)"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_WHITE

    # =========================================================================
    # SLIDE 7: LANGKAH 6: DISKRETISASI DATA & KONSEP HIERARKI
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "Langkah 6: Diskretisasi Data & Pembentukan Konsep Hierarki")

    add_card(slide7, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    tb = slide7.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "1. Kelulusan Pendidikan Dasar"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)

    bins1 = [
        ("Kelulusan Rendah (<96%):", "Tantangan efisiensi penuntasan (contoh: China 92,47%)."),
        ("Kelulusan Sedang (96% - 99%):", "Penuntasan stabil tingkat menengah (contoh: India 97,55%)."),
        ("Kelulusan Tinggi (>99%):", "Penuntasan universal penuh & kelulusan siswa lintas usia (contoh: Jerman 102,5%, Jepang 102,0%, Indonesia 102,0%, AS 100,6%)."),
        ("💡 Catatan Standar UNESCO UIS:", "Indikator Gross Intake Ratio to Last Grade mencatat kelulusan siswa lintas usia (over-age graduates), sehingga angka >100% merupakan hal valid.")
    ]
    for t, d in bins1:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK_TEXT
        run = p.add_run()
        run.text = d
        run.font.bold = False
        p.space_after = Pt(6)

    add_card(slide7, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
    tb2 = slide7.shapes.add_textbox(Inches(7.2), Inches(1.7), Inches(5.0), Inches(4.9))
    tf = tb2.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "2. Rasio Murid per Guru (Beban Kelas)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)

    bins2 = [
        ("Rasio Ideal (<15 murid/guru):", "Kapasitas pengajaran optimal negara maju (contoh: Jerman 12,22; AS 14,54)."),
        ("Rasio Moderat (15 - 20 murid/guru):", "Kapasitas standar seimbang (contoh: Indonesia 16,56; Jepang 16,45; China 16,29; UK 17,39)."),
        ("Rasio Padat (>20 murid/guru):", "Beban pengajaran kelas tinggi (contoh: India 31,49; Brasil 20,92)."),
        ("💡 Posisi Benchmarking Indonesia:", "Indonesia menempati Kategori Rasio Moderat (~16 murid per guru), jauh lebih optimal dibanding negara berkembang seperti India.")
    ]
    for t, d in bins2:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK_TEXT
        run = p.add_run()
        run.text = d
        run.font.bold = False
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 8: EVALUASI KUANTITATIF SEBELUM VS SESUDAH
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "Evaluasi Kuantitatif: Sebelum vs. Sesudah Preprocessing")

    # Left: Bar chart image
    add_card(slide8, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.1), bg_color=COLOR_WHITE)
    if os.path.exists('before_after_comparison.png'):
        slide8.shapes.add_picture('before_after_comparison.png', Inches(0.95), Inches(1.65), width=Inches(5.3))

    # Right: Case study image
    add_card(slide8, Inches(6.9), Inches(1.5), Inches(5.6), Inches(4.1), bg_color=COLOR_WHITE)
    if os.path.exists('before_after_imputation_case_study.png'):
        slide8.shapes.add_picture('before_after_imputation_case_study.png', Inches(7.05), Inches(1.65), width=Inches(5.3))

    # Bottom summary card
    add_card(slide8, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.3), bg_color=COLOR_BG_CARD, border_color=COLOR_SECONDARY)
    tb = slide8.shapes.add_textbox(Inches(1.0), Inches(5.85), Inches(11.3), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Ringkasan Hasil Evaluasi Kualitas Pipeline Preprocessing:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    p = tf.add_paragraph()
    p.text = "• Missing Rate: 86,10% (Data Mentah) -> 46,54% (Panel Awal) -> 0,00% (Data Bersih via 3-Tier Imputation)\n• Efisiensi Komputasi: Ukuran berkas terpangkas dari 326,4 MB menjadi 45,8 KB (>99,9% penghematan penyimpanan & RAM 0,04 MB)\n• Keberhasilan Imputasi Deret Waktu: Menyambungkan gap awal/akhir Indonesia dan gap tengah India tanpa mengubah arah akselerasi."
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_DARK_TEXT

    # =========================================================================
    # SLIDE 9: TEMUAN ANALITIS & BENCHMARKING G20
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "Temuan Analitis & Wawasan Pendidikan Negara G20")

    # Left: Heatmap
    add_card(slide9, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.2), bg_color=COLOR_WHITE)
    if os.path.exists('temporal_heatmap_g20.png'):
        slide9.shapes.add_picture('temporal_heatmap_g20.png', Inches(0.95), Inches(1.65), width=Inches(5.3))

    # Right: Gender parity
    add_card(slide9, Inches(6.9), Inches(1.5), Inches(5.6), Inches(4.2), bg_color=COLOR_WHITE)
    if os.path.exists('gender_parity_analysis.png'):
        slide9.shapes.add_picture('gender_parity_analysis.png', Inches(7.05), Inches(1.65), width=Inches(5.3))

    # Bottom summary card
    add_card(slide9, Inches(0.8), Inches(5.85), Inches(11.7), Inches(1.2), bg_color=COLOR_BG_CARD, border_color=COLOR_CARD_BORDER)
    tb = slide9.shapes.add_textbox(Inches(1.0), Inches(5.9), Inches(11.3), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Wawasan Analitis Kebijakan Pendidikan:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    p = tf.add_paragraph()
    p.text = "1. Konvergensi Partisipasi Global (Heatmap): Negara berkembang G20 (India & Turki) mengalami peningkatan pesat menuju akses universal (>95%).\n2. Paritas Gender Sempurna (1:1 Ratio): Seluruh observasi menempel rapat pada garis y=x, membuktikan tidak ada disparitas gender akses pendidikan di G20.\n3. Posisi Benchmarking Indonesia: NER stabil di kisaran 96%–98%, berada di barisan teratas negara berkembang dan mendekati negara OECD."
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_DARK_TEXT

    # =========================================================================
    # SLIDE 10: KESIMPULAN TENTANG PROSES PREPROCESSING
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide10, "Kesimpulan Proses Data Preprocessing")

    c_points = [
        ("1. Integrasi & Resolusi Format:", "Berhasil menggabungkan 5 berkas mentah ke dalam skema bintang (Star Schema) dan membersihkan inkonsistensi representasi teks 'YR2000' menjadi integer."),
        ("2. Penuntasan Missing Values (0,00%):", "Penerapan Imputasi Cerdas 3-Tier (interpolasi temporal, median kelompok pendapatan, dan median global) sukses menuntaskan kekosongan data tanpa merusak pola tren."),
        ("3. Efisiensi Komputasi Masif (>99,9%):", "Transformasi Tidy Panel mereduksi ukuran berkas dari 326,4 MB menjadi 45,8 KB serta alokasi RAM dari 740 MB menjadi 0,04 MB."),
        ("4. Retensi Informasi Optimal via PCA:", "Mereduksi 6 dimensi indikator menjadi 2 Komponen Utama (PC1 & PC2) dengan tetap mempertahankan 84,47% variansi informasi aslinya."),
        ("5. Dataset Bersih Siap Pakai (ML Ready):", "Menghasilkan artefak CSV bersih 'EdStats_Cleaned_G20_2000_2017.csv' (342 observasi x 13 fitur) yang terstruktur dan siap untuk pemodelan data mining.")
    ]

    add_card(slide10, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3), bg_color=COLOR_BG_CARD, border_color=COLOR_SECONDARY)
    tb = slide10.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(11.1), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Pencapaian Utama Rekayasa Data (Pipeline Summary):"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    for t, d in c_points:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_PRIMARY
        run = p.add_run()
        run.text = d
        run.font.bold = False
        run.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 11: CLOSING / Q&A
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_slide_layout)
    add_card(slide11, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=COLOR_PRIMARY, border_color=None)
    add_card(slide11, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1), bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER)

    tb = slide11.shapes.add_textbox(Inches(1.6), Inches(1.8), Inches(10.1), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "TERIMA KASIH"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Sesi Tanya Jawab (Q & A)"
    p.font.size = Pt(18)
    p.font.color.rgb = COLOR_SECONDARY
    p.space_after = Pt(20)

    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Rekayasa Data (Data Engineering)\nDepartemen Teknik Elektro dan Teknologi Informasi, Fakultas Teknik UGM"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK_TEXT
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Tautan Sumber Terbuka: Kaggle EdStats & World Bank Official Data Catalog"
    p.font.size = Pt(10.5)
    p.font.italic = True
    p.font.color.rgb = COLOR_MUTED_TEXT

    prs.save(output_pptx)
    print(f"File presentasi PPTX berhasil dibuat: {output_pptx} (Ukuran: {os.path.getsize(output_pptx)/1024:.1f} KB)")

if __name__ == '__main__':
    create_deck()
